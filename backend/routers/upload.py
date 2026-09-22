from fastapi import APIRouter, UploadFile, File, HTTPException
import pdfplumber
from pptx import Presentation
import tempfile
import os

router = APIRouter()

MAX_FILE_SIZE = 20 * 1024 * 1024  # 20MB


def extract_pdf(file_path: str):
    slides = []
    with pdfplumber.open(file_path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text()
            if text and text.strip():
                slides.append({"slide_number": i + 1, "text": text.strip()})
    return slides

def extract_pptx(file_path: str):
    slides = []
    prs = Presentation(file_path)
    for i, slide in enumerate(prs.slides):
        text = " ".join(
            shape.text for shape in slide.shapes if hasattr(shape, "text")
        ).strip()
        if text:
            slides.append({"slide_number": i + 1, "text": text})
    return slides

@router.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in [".pdf", ".pptx"]:
        raise HTTPException(status_code=400, detail="Only PDF and PPTX supported")

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="File too large (max 20MB)")

    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        tmp.write(content)
        tmp_path = tmp.name

    try:
        if ext == ".pdf":
            slides = extract_pdf(tmp_path)
        else:
            slides = extract_pptx(tmp_path)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read file, please try again")
    finally:
        os.unlink(tmp_path)

    if not slides:
        raise HTTPException(status_code=422, detail="No readable text found in slides")

    return {"total_slides": len(slides), "slides": slides}